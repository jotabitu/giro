"""T-027 · Importação de veículos e clientes por planilha (docs/api/importacao.md, IMP-02,
IMP-03, IMP-05).

Roda na `empresa_a` com a Filial Centro e a tabela de preços da Moto Rápida, sem veículos, com
o relógio em 07 out 2026, 10:00. O caso do seed: 86 linhas, 83 prontas e 3 com erro (as da tela
`Onboarding-Importar`); corrigidas as 3, 86 veículos.
"""

import csv
import io
import zipfile
from typing import Any

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django_tenants.utils import schema_context
from openpyxl import Workbook, load_workbook

from apps.auditoria.models import EventoAuditoria
from apps.clientes.models import Cliente
from apps.contas.models import Usuario
from apps.importacao import selectors, services
from apps.importacao.models import Importacao, LinhaImportacao
from apps.notificacoes.models import Notificacao
from apps.plataforma.models import Empresa, UsoEmpresa
from apps.plataforma.services import abrir_acesso_suporte
from apps.veiculos.models import HistoricoStatusVeiculo, Veiculo
from comum import arquivos, relogio

from .conftest import cliente as cliente_http
from .frota import criar_categorias, criar_veiculo
from .test_t012_filiais import HOST, MotoRapida

URL = "/api/importacoes/"
COLUNAS_VEICULO = [
    "placa",
    "chassi",
    "marca",
    "modelo",
    "ano_fabricacao",
    "ano_modelo",
    "cor",
    "km",
    "categoria",
]
# Quantos de cada categoria (a frota da Moto Rápida: 42 + 32 + 6 + 6 = 86).
CATEGORIAS = [
    ("Urbana 110–125", 42),
    ("Street 150–160", 32),
    ("Scooter premium", 6),
    ("Carro compacto", 6),
]
COLUNAS_CLIENTE = ["nome", "documento", "celular", "email", "nascimento", "cnh_numero"]
JULIANA = ["Juliana Alves", "052.218.330-18", "(85) 99812-3344", "", "14/03/1998", "04512345678"]
MARCOS = ["Marcos Lima", "529.982.247-25", "(85) 98877-1122", "", "", ""]


@pytest.fixture
def moto(empresa_a, planos, relogio_jwt) -> MotoRapida:
    moto = MotoRapida(empresa_a, planos["Profissional"])
    moto.filial = moto.centro()
    moto.categorias = criar_categorias(moto.pessoas["Mariana"])
    return moto


def _xlsx(cabecalho: list[str], linhas: list[list[Any]]) -> bytes:
    livro = Workbook()
    folha = livro.worksheets[0]
    folha.append(cabecalho)
    for linha in linhas:
        folha.append(linha)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _frota(quantos: int = 86) -> list[list[Any]]:
    """A planilha da Moto Rápida. Linha i da lista = linha i + 2 da planilha."""
    categorias = [nome for nome, n in CATEGORIAS for _ in range(n)]
    return [
        [
            f"QAA{i:04d}",
            f"9C2KC2210PR{i:06d}",
            "Honda",
            "CG 160 Fan",
            2024,
            2024,
            "Preta",
            1000 + i,
            categorias[i],
        ]
        for i in range(quantos)
    ]


def _frota_com_os_3_erros() -> list[list[Any]]:
    linhas = _frota()
    linhas[12][0] = "QXA4B2"  # linha 14: placa incompleta
    linhas[35][1] = linhas[10][1]  # linha 37: o chassi da linha 12
    linhas[50][5] = None  # linha 52: ano do modelo em branco
    return linhas


def _enviar(api, conteudo: bytes, nome: str = "frota-moto-rapida.xlsx", tipo="veiculos", **mais):
    arquivo = SimpleUploadedFile(nome, conteudo)
    return api.post(URL, {"tipo": tipo, "arquivo": arquivo, **mais}, format="multipart")


def _rodando(capturar, funcao, *args, **kwargs):
    """As tarefas são agendadas no commit; aqui elas rodam na hora (CELERY_TASK_ALWAYS_EAGER)."""
    tenant = connection.tenant
    with capturar(execute=True):
        resposta = funcao(*args, **kwargs)
    connection.set_tenant(tenant)  # o fim da tarefa devolve a conexão ao public (comum.worker)
    return resposta


def _linhas(api, importacao_id: int, status: str = "") -> list[dict[str, Any]]:
    consulta = f"?status={status}&tamanho=100" if status else "?tamanho=100"
    resposta = api.get(f"{URL}{importacao_id}/linhas/{consulta}")
    assert resposta.status_code == 200, resposta.json()
    return list(resposta.json()["resultados"])


def _importar(api, capturar, importacao_id: int):
    return _rodando(capturar, api.post, f"{URL}{importacao_id}/importar/", {}, format="json")


def _suporte(moto: MotoRapida, permissao: str):
    with schema_context("public"):
        diego = Usuario.objects.filter(email="diego@girolocadoras.com.br").first()
        if diego is None:
            diego = Usuario.objects.create_user(
                "diego@girolocadoras.com.br",
                "senha-forte-123",
                nome="Diego Santos",
                funcao_plataforma="suporte",
            )
    _, token = abrir_acesso_suporte(
        Empresa.objects.get(pk=moto.empresa.pk),
        diego,
        motivo="Chamado #1042",
        duracao_minutos=30,
        permissao=permissao,
    )
    connection.set_tenant(moto.empresa)
    return cliente_http(HOST, str(token))


# --- IMP-02: o caso da tela ---


def test_imp_02_83_prontos_3_com_erro_e_86_depois_de_corrigir(
    moto, django_capture_on_commit_callbacks
):
    api = moto.api()
    enviada = _rodando(
        django_capture_on_commit_callbacks,
        _enviar,
        api,
        _xlsx(COLUNAS_VEICULO, _frota_com_os_3_erros()),
    )
    assert enviada.status_code == 202, enviada.json()
    assert enviada.json()["status"] == "validando"
    imp = enviada.json()["id"]

    corpo = api.get(f"{URL}{imp}/").json()
    assert corpo["status"] == "validada"
    assert corpo["arquivo"] == "frota-moto-rapida.xlsx"
    assert corpo["filial_padrao"] == {"id": moto.filial.pk, "nome": "Filial Centro"}
    assert corpo["contagens"] == {
        "linhas": 86,
        "prontas": 83,
        "com_erro": 3,
        "importadas": 0,
        "acima_do_limite": 0,
    }
    assert corpo["plano"] == {"usado": 0, "limite": 150, "cabem": 150}
    assert corpo["enviada_por"] == "Mariana"
    assert corpo["apagar_em"] == "2026-11-06"

    com_erro = _linhas(api, imp, "com_erro")
    assert [(li["numero"], li["erros"]) for li in com_erro] == [
        (14, [{"campo": "placa", "codigo": "VEI-01", "mensagem": "Placa QXA4B2 está incompleta"}]),
        (37, [{"campo": "chassi", "codigo": "VEI-01", "mensagem": "Chassi repetido na linha 12"}]),
        (
            52,
            [{"campo": "ano_modelo", "codigo": "validacao", "mensagem": "Ano do modelo em branco"}],
        ),
    ]
    assert Veiculo.objects.count() == 0  # validar não grava nada

    importada = _importar(api, django_capture_on_commit_callbacks, imp)
    assert importada.status_code == 202, importada.json()
    corpo = api.get(f"{URL}{imp}/").json()
    assert corpo["status"] == "concluida"
    assert corpo["contagens"]["importadas"] == 83
    assert corpo["contagens"]["com_erro"] == 3
    assert Veiculo.objects.count() == 83
    assert UsoEmpresa.objects.get(empresa=moto.empresa).veiculos == 83

    # "Corrigir aqui": as três, revalidadas na hora.
    for numero, dados in (
        (14, {"placa": "QXA4B21"}),
        (37, {"chassi": "9C2KC2210PR999999"}),
        (52, {"ano_modelo": "2024"}),
    ):
        resposta = api.patch(f"{URL}{imp}/linhas/{numero}/", {"dados": dados}, format="json")
        assert resposta.status_code == 200, resposta.json()
        assert resposta.json()["status"] == "pronta", resposta.json()
        assert resposta.json()["erros"] == []
    assert api.get(f"{URL}{imp}/").json()["contagens"]["prontas"] == 3

    _importar(api, django_capture_on_commit_callbacks, imp)
    corpo = api.get(f"{URL}{imp}/").json()
    assert corpo["contagens"] == {
        "linhas": 86,
        "prontas": 0,
        "com_erro": 0,
        "importadas": 86,
        "acima_do_limite": 0,
    }
    assert Veiculo.objects.count() == 86
    qxa = Veiculo.objects.get(placa="QXA4B21")
    assert qxa.status == "disponivel"
    assert qxa.filial_id == moto.filial.pk
    assert qxa.categoria.nome == "Urbana 110–125"
    historico = HistoricoStatusVeiculo.objects.get(veiculo=qxa)
    assert (historico.origem_tipo, historico.origem_id) == ("importacao", imp)
    linha_14 = _linhas(api, imp, "importada")[12]
    assert linha_14["numero"] == 14
    assert linha_14["objeto"] == {"tipo": "veiculo", "id": qxa.pk}


def test_imp_02_correcao_que_ainda_tem_erro_continua_com_erro(
    moto, django_capture_on_commit_callbacks
):
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks,
        _enviar,
        api,
        _xlsx(COLUNAS_VEICULO, _frota_com_os_3_erros()),
    ).json()["id"]
    resposta = api.patch(f"{URL}{imp}/linhas/14/", {"dados": {"placa": "QAA0002"}}, format="json")
    assert resposta.json()["status"] == "com_erro"
    assert resposta.json()["erros"] == [
        {"campo": "placa", "codigo": "VEI-01", "mensagem": "Placa repetida na linha 4"}
    ]
    # Linha importada não aceita correção.
    _importar(api, django_capture_on_commit_callbacks, imp)
    fechada = api.patch(f"{URL}{imp}/linhas/2/", {"dados": {"cor": "Azul"}}, format="json")
    assert fechada.status_code == 409
    assert fechada.json()["codigo"] == "linha_fechada"


def test_imp_02_ja_cadastrado_e_erro_nao_atualizacao(moto, django_capture_on_commit_callbacks):
    existente = criar_veiculo(1, moto.categorias["Urbana 110–125"], moto.filial, cor="Vermelha")
    linhas = _frota(3)
    linhas[0][0] = existente.placa
    linhas[1][1] = existente.chassi
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks, _enviar, api, _xlsx(COLUNAS_VEICULO, linhas)
    ).json()["id"]
    erros = [li["erros"] for li in _linhas(api, imp, "com_erro")]
    assert erros == [
        [
            {
                "campo": "placa",
                "codigo": "VEI-01",
                "mensagem": "Placa GEN0001 já cadastrada",
            }
        ],
        [
            {
                "campo": "chassi",
                "codigo": "VEI-01",
                "mensagem": f"Chassi {existente.chassi} já cadastrado",
            }
        ],
    ]
    _importar(api, django_capture_on_commit_callbacks, imp)
    existente.refresh_from_db()
    assert existente.cor == "Vermelha"
    assert Veiculo.objects.count() == 2


def test_imp_02_categoria_filial_e_numeros(moto, django_capture_on_commit_callbacks):
    cabecalho = [*COLUNAS_VEICULO, "filial", "valor_compra", "data_compra"]
    linhas = [
        [
            "QAA0001",
            "9C2KC2210PR000001",
            "Honda",
            "CG",
            2024,
            2024,
            "Preta",
            "1.200",
            "street 150-160",
            "filial centro",
            "13.900,00",
            "10/01/2025",
        ],
        [
            "QAA0002",
            "9C2KC2210PR000002",
            "Honda",
            "CG",
            "dois mil",
            2024,
            "Preta",
            "abc",
            "Inexistente",
            "",
            "",
            "31/02/2025",
        ],
        [
            "QAA0003",
            "9C2KC2210PR000003",
            "Honda",
            "CG",
            2024,
            2024,
            "Preta",
            10,
            "Street 150–160",
            "Filial Norte",
            "13900.00",
            "",
        ],
    ]
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks, _enviar, api, _xlsx(cabecalho, linhas)
    ).json()["id"]
    todas = _linhas(api, imp)
    assert todas[0]["status"] == "pronta", todas[0]
    assert [e["mensagem"] for e in todas[1]["erros"]] == [
        "Ano de fabricação precisa ser um número",
        "Km precisa ser um número",
        "Categoria Inexistente não existe ou está inativa",
        "Data de compra inválida",
    ]
    assert [e["mensagem"] for e in todas[2]["erros"]] == [
        "Filial Filial Norte não existe ou está inativa"
    ]
    _importar(api, django_capture_on_commit_callbacks, imp)
    gravado = Veiculo.objects.get(placa="QAA0001")
    assert (gravado.km_atual, gravado.valor_compra_centavos) == (1200, 1390000)
    assert str(gravado.data_compra) == "2025-01-10"


# --- IMP-03: limite do plano ---


def test_imp_03_importa_ate_o_limite_e_avisa(moto, planos, django_capture_on_commit_callbacks):
    moto.empresa.plano = planos["Essencial"]
    moto.empresa.save(update_fields=["plano"])
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks, _enviar, api, _xlsx(COLUNAS_VEICULO, _frota(42))
    ).json()["id"]
    assert api.get(f"{URL}{imp}/").json()["plano"] == {"usado": 0, "limite": 30, "cabem": 30}

    _importar(api, django_capture_on_commit_callbacks, imp)
    corpo = api.get(f"{URL}{imp}/").json()
    assert corpo["contagens"]["importadas"] == 30
    assert corpo["contagens"]["acima_do_limite"] == 12
    assert corpo["aviso"] == {
        "codigo": "IMP-03",
        "mensagem": "Importamos 30 veículos, o limite do plano Essencial. 12 ficaram de fora: "
        "mude de plano para importar o resto.",
    }
    assert Veiculo.objects.count() == 30
    fora = _linhas(api, imp, "acima_do_limite")
    assert [li["numero"] for li in fora] == list(range(32, 44))  # as últimas, na ordem

    # Mudando de plano, importar de novo pega as que ficaram de fora.
    moto.empresa.plano = planos["Profissional"]
    moto.empresa.save(update_fields=["plano"])
    _importar(api, django_capture_on_commit_callbacks, imp)
    corpo = api.get(f"{URL}{imp}/").json()
    assert corpo["contagens"]["importadas"] == 42
    assert corpo["aviso"] is None
    assert Veiculo.objects.count() == 42


def test_imp_05_conclusao_marca_o_passo_audita_e_avisa(moto, django_capture_on_commit_callbacks):
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks, _enviar, api, _xlsx(COLUNAS_VEICULO, _frota(2))
    ).json()["id"]
    _importar(api, django_capture_on_commit_callbacks, imp)
    assert "importar" in UsoEmpresa.objects.get(empresa=moto.empresa).passos_implantacao
    concluir = EventoAuditoria.objects.get(acao="importacao.concluir")
    assert (concluir.categoria, concluir.ator_nome, concluir.ator_tipo) == (
        "configuracoes",
        "Mariana",
        "usuario",
    )
    assert concluir.depois["importadas"] == 2
    cadastros = EventoAuditoria.objects.filter(acao="veiculo.cadastrar")
    assert {e.ator_nome for e in cadastros} == {"Mariana"}
    aviso = Notificacao.objects.get(tipo="importacao_concluida")
    assert aviso.destinatario_id == moto.pessoas["Mariana"].pk
    assert aviso.titulo == "Importação concluída: 2 veículos importados"


# --- Clientes ---


def test_cli_01_clientes_validados_cifrados_e_incompletos(moto, django_capture_on_commit_callbacks):
    linhas = [
        JULIANA,
        ["Pedro Souza", "123.456.789-00", "(85) 98765-4321", "pedro@", "40/01/1990", ""],
        ["Juliana de novo", "05221833018", "(85) 99812-0000", "", "", ""],
        MARCOS,
    ]
    api = moto.api()
    enviada = _rodando(
        django_capture_on_commit_callbacks,
        _enviar,
        api,
        _xlsx(COLUNAS_CLIENTE, linhas),
        "clientes.xlsx",
        "clientes",
    )
    imp = enviada.json()["id"]
    corpo = api.get(f"{URL}{imp}/").json()
    assert corpo["plano"] is None and corpo["filial_padrao"] is None
    todas = _linhas(api, imp)
    assert [li["status"] for li in todas] == ["pronta", "com_erro", "com_erro", "pronta"]
    assert todas[1]["erros"] == [
        {"campo": "nascimento", "codigo": "validacao", "mensagem": "Nascimento inválida"},
        {"campo": "documento", "codigo": "CLI-01", "mensagem": "CPF inválido"},
        {"campo": "email", "codigo": "validacao", "mensagem": "E-mail inválido"},
    ]
    assert todas[2]["erros"] == [
        {"campo": "documento", "codigo": "CLI-01", "mensagem": "CPF já cadastrado (linha 2)"}
    ]
    # O administrador vê a linha inteira.
    assert todas[0]["dados"]["documento"] == "052.218.330-18"
    # Os erros não carregam dado pessoal; as linhas ficam cifradas no banco.
    texto_dos_erros = str([li["erros"] for li in todas])
    for aberto in ("123.456.789-00", "pedro@", "05221833018"):
        assert aberto not in texto_dos_erros
    guardada = LinhaImportacao.objects.get(importacao_id=imp, numero=2)
    assert guardada.dados == {} and "Juliana" not in guardada.dados_cifrados

    _importar(api, django_capture_on_commit_callbacks, imp)
    juliana = Cliente.objects.get(nome="Juliana Alves")
    assert juliana.status == "incompleto"  # sem fotos nem consentimento (CLI-03)
    assert juliana.documento_claro() == "05221833018"
    assert juliana.cnh_numero_claro() == "04512345678"
    assert Cliente.objects.count() == 2


def test_importacao_arquivo_de_clientes_apagado_apos_validar(
    moto, django_capture_on_commit_callbacks
):
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks,
        _enviar,
        api,
        _xlsx(COLUNAS_CLIENTE, [JULIANA]),
        "clientes.xlsx",
        "clientes",
    ).json()["id"]
    importacao = Importacao.objects.get(pk=imp)
    assert importacao.arquivo_caminho == ""
    assert not [c for c in arquivos.guardados if "/importacao/" in c and f"/{imp}/" in c]
    # O de veículos fica até `apagar_em`, com nome aleatório sob o schema (AUD-05).
    veiculos = _rodando(
        django_capture_on_commit_callbacks, _enviar, api, _xlsx(COLUNAS_VEICULO, _frota(1))
    ).json()["id"]
    caminho = Importacao.objects.get(pk=veiculos).arquivo_caminho
    assert caminho.startswith(f"{moto.empresa.schema_name}/importacao/importacao/{veiculos}/")
    assert "frota" not in caminho and caminho in arquivos.guardados


def test_aud_09_suporte_ve_linha_de_cliente_resumida(moto, django_capture_on_commit_callbacks):
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks,
        _enviar,
        api,
        _xlsx(COLUNAS_CLIENTE, [JULIANA]),
        "clientes.xlsx",
        "clientes",
    ).json()["id"]
    for permissao in ("ver", "ver_editar"):
        suporte = _suporte(moto, permissao)
        resposta = suporte.get(f"{URL}{imp}/linhas/")
        assert resposta.status_code == 200, resposta.json()
        assert resposta.json()["resultados"][0]["dados"] == {
            "nome": "Juliana Alves",
            "tipo": "",
            "documento_mascarado": "***.218.330-**",
            "celular": "(85) 99812-3344",
        }
        texto = resposta.content.decode()
        for aberto in ("052.218.330-18", "04512345678", "14/03/1998"):
            assert aberto not in texto
    # "Ver e editar" corrige, e a resposta continua resumida.
    corrigida = suporte.patch(
        f"{URL}{imp}/linhas/2/", {"dados": {"celular": "85999990000"}}, format="json"
    )
    assert corrigida.status_code == 200, corrigida.json()
    assert "documento" not in corrigida.json()["dados"]
    # "Só ver" não corrige.
    so_ver = _suporte(moto, "ver")
    assert (
        so_ver.patch(f"{URL}{imp}/linhas/2/", {"dados": {"nome": "X"}}, format="json").status_code
        == 403
    )


def test_aud_03_importacao_pelo_suporte_audita_ator_suporte(
    moto, django_capture_on_commit_callbacks
):
    suporte = _suporte(moto, "ver_editar")
    enviada = _rodando(
        django_capture_on_commit_callbacks, _enviar, suporte, _xlsx(COLUNAS_VEICULO, _frota(2))
    )
    assert enviada.status_code == 202, enviada.json()
    assert enviada.json()["enviada_por"] == "Suporte Giro"
    _importar(suporte, django_capture_on_commit_callbacks, enviada.json()["id"])
    assert Veiculo.objects.count() == 2
    for acao in ("importacao.enviar", "importacao.concluir", "veiculo.cadastrar"):
        eventos = EventoAuditoria.objects.filter(acao=acao)
        assert eventos.exists(), acao
        assert {e.ator_tipo for e in eventos} == {"suporte"}, acao
        assert all(e.acesso_suporte_id for e in eventos), acao
    # Sem pessoa da empresa que pediu: o aviso vai para os administradores.
    aviso = Notificacao.objects.get(tipo="importacao_concluida")
    assert aviso.destinatario_id == moto.pessoas["Mariana"].pk


# --- Arquivo (segurança, revisão T-044) ---


def test_importacao_xlsx_com_entidades_recusado(moto):
    original = _xlsx(COLUNAS_VEICULO, _frota(1))
    folha = (
        '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa">'
        '<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&b;</t></is></c></row>'
        "</sheetData></worksheet>"
    )
    saida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as entrada, zipfile.ZipFile(saida, "w") as novo:
        for item in entrada.infolist():
            dados = entrada.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                dados = folha.encode()
            novo.writestr(item, dados)
    resposta = _enviar(moto.api(), saida.getvalue())
    assert resposta.status_code == 400
    assert resposta.json()["campos"] == {
        "arquivo": ["Não conseguimos ler a planilha. Confira o arquivo."]
    }
    assert not Importacao.objects.exists()


def test_importacao_zip_acima_de_50mb_recusado(moto):
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as pacote:
        pacote.writestr("[Content_Types].xml", "<Types/>")
        pacote.writestr("xl/bomba.xml", b"0" * (51 * 1024 * 1024))
    assert len(saida.getvalue()) < 5 * 1024 * 1024
    resposta = _enviar(moto.api(), saida.getvalue())
    assert resposta.status_code == 400
    assert resposta.json()["campos"] == {"arquivo": ["A planilha é grande demais para importar."]}


def test_importacao_conteudo_extensao_tamanho_e_cabecalho(moto):
    api = moto.api()
    for conteudo, nome, mensagem in (
        (b"placa;chassi", "frota.pdf", "Envie a planilha em .xlsx ou .csv."),
        (b"", "frota.csv", "O arquivo está vazio."),
        (b"isto nao e zip", "frota.xlsx", "O arquivo não é uma planilha .xlsx."),
        (b"x" * (5 * 1024 * 1024 + 1), "frota.csv", "O arquivo passa de 5 MB."),
        (
            "placa;chassi\n".encode("latin-1") + "ç".encode("latin-1"),
            "frota.csv",
            "Salve o .csv em UTF-8 e envie de novo.",
        ),
    ):
        resposta = _enviar(api, conteudo, nome)
        assert resposta.status_code == 400, nome
        assert resposta.json()["campos"] == {"arquivo": [mensagem]}, nome
    sem_colunas = _enviar(api, _xlsx(["Placa", "Chassi", "Marca"], [["QAA0001", "x", "Honda"]]))
    assert sem_colunas.status_code == 400
    assert sem_colunas.json()["codigo"] == "validacao"
    assert sem_colunas.json()["detalhes"]["colunas_faltando"] == [
        "modelo",
        "ano_fabricacao",
        "ano_modelo",
        "km",
        "categoria",
    ]
    linhas = [[f"QAA{i:04d}", "", "", "", "", "", "", "", ""] for i in range(5001)]
    grande = _enviar(api, _xlsx(COLUNAS_VEICULO, linhas))
    assert grande.json()["campos"] == {"arquivo": ["A planilha passa de 5.000 linhas."]}


def test_importacao_csv_com_bom_ponto_e_virgula_e_cabecalho_com_acento(
    moto, django_capture_on_commit_callbacks
):
    texto = io.StringIO()
    escritor = csv.writer(texto, delimiter=";")
    escritor.writerow(
        [
            "Placa",
            "Chassi",
            "Marca",
            "Modelo",
            "Ano Fabricação",
            "ANO-MODELO",
            "Cor",
            "KM",
            "Categoria",
            "Observação",
        ]
    )
    escritor.writerow(
        [
            "qaa-0001",
            "9c2kc2210pr000001",
            "Honda",
            "CG",
            "2024",
            "2024",
            "Preta",
            "10",
            "Urbana 110–125",
            "nada",
        ]
    )
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks,
        _enviar,
        api,
        ("﻿" + texto.getvalue()).encode("utf-8"),
        "frota.csv",
    ).json()["id"]
    todas = _linhas(api, imp)
    assert todas[0]["status"] == "pronta"
    assert "observacao" not in todas[0]["dados"]  # coluna não reconhecida sai
    assert todas[0]["dados"]["placa"] == "qaa-0001"
    # No .csv, a célula acima de 1.000 caracteres (`csv.field_size_limit`) recusa o arquivo.
    escritor.writerow(
        ["QAA0002", "x", "Honda", "CG", "2024", "2024", "Preta", "10", "U", "y" * 1001]
    )
    longa = _enviar(api, texto.getvalue().encode("utf-8"), "frota.csv")
    assert longa.status_code == 400
    assert longa.json()["campos"] == {
        "arquivo": ["A linha 3 tem uma célula com mais de 1.000 caracteres."]
    }


def test_importacao_celula_longa_no_xlsx_deixa_a_linha_com_erro(
    moto, django_capture_on_commit_callbacks
):
    linhas = _frota(1)
    linhas[0][6] = "y" * 1001
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks, _enviar, api, _xlsx(COLUNAS_VEICULO, linhas)
    ).json()["id"]
    linha = _linhas(api, imp)[0]
    assert linha["erros"] == [
        {"campo": "cor", "codigo": "validacao", "mensagem": "Célula com mais de 1.000 caracteres"}
    ]


# --- Fluxo ---


def test_importacao_modelo_para_baixar(moto):
    for tipo, primeira in (("veiculos", "placa"), ("clientes", "tipo")):
        resposta = moto.api().get(f"{URL}modelo/?tipo={tipo}")
        assert resposta.status_code == 200
        livro = load_workbook(io.BytesIO(resposta.content))
        assert livro.sheetnames[1] == "Instruções"
        folha = livro.worksheets[0]
        assert folha["A1"].value == primeira
        assert folha.max_row == 2  # cabeçalho e a linha de exemplo
    assert moto.api().get(f"{URL}modelo/?tipo=outro").status_code == 400


def test_importacao_so_administrador(moto):
    for nome in ("Rodrigo", "Bruna", "João", "Lúcia"):
        assert moto.api(nome).get(URL).status_code == 403, nome
        assert _enviar(moto.api(nome), _xlsx(COLUNAS_VEICULO, _frota(1))).status_code == 403


def test_importacao_em_andamento_cancelar_e_nada_para_importar(
    moto, django_capture_on_commit_callbacks
):
    api = moto.api()
    parada = Importacao.objects.create(
        tipo="veiculos",
        arquivo_nome="outra.xlsx",
        enviada_em=relogio.agora(),
        apagar_em=relogio.hoje(),
    )
    ocupada = _enviar(api, _xlsx(COLUNAS_VEICULO, _frota(1)))
    assert ocupada.status_code == 409
    assert ocupada.json()["codigo"] == "importacao_em_andamento"
    # Clientes não esperam a de veículos.
    livre = _enviar(api, _xlsx(COLUNAS_CLIENTE, [JULIANA]), "c.xlsx", "clientes")
    assert livre.status_code == 202
    parada.delete()

    imp = _rodando(
        django_capture_on_commit_callbacks, _enviar, api, _xlsx(COLUNAS_VEICULO, _frota(1))
    ).json()["id"]
    ignorada = api.post(f"{URL}{imp}/linhas/2/ignorar/", {}, format="json")
    assert ignorada.json()["status"] == "ignorada"
    nada = api.post(f"{URL}{imp}/importar/", {}, format="json")
    assert nada.status_code == 422
    assert nada.json()["codigo"] == "nada_para_importar"
    cancelada = api.post(f"{URL}{imp}/cancelar/", {}, format="json")
    assert cancelada.json()["status"] == "cancelada"
    assert api.post(f"{URL}{imp}/cancelar/", {}, format="json").status_code == 409
    assert Veiculo.objects.count() == 0


def test_importacao_filial_obrigatoria_com_mais_de_uma(moto, django_capture_on_commit_callbacks):
    from .test_t012_filiais import _filial

    assert moto.criar(_filial("Filial Caucaia")).status_code == 201
    api = moto.api()
    sem = _enviar(api, _xlsx(COLUNAS_VEICULO, _frota(1)))
    assert sem.status_code == 400
    assert sem.json()["campos"] == {"filial_id": ["Escolha a filial dos veículos."]}
    com = _enviar(api, _xlsx(COLUNAS_VEICULO, _frota(1)), filial_id=moto.filial.pk)
    assert com.status_code == 202


def test_importacao_limpeza_depois_de_30_dias(moto, relogio, django_capture_on_commit_callbacks):
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks, _enviar, api, _xlsx(COLUNAS_VEICULO, _frota(2))
    ).json()["id"]
    caminho = Importacao.objects.get(pk=imp).arquivo_caminho
    assert services.limpar_vencidas() == 0
    relogio.em("2026-11-06 10:00")
    assert services.limpar_vencidas() == 1
    assert not Importacao.objects.filter(pk=imp).exists()
    assert not LinhaImportacao.objects.filter(importacao_id=imp).exists()
    assert caminho not in arquivos.guardados


def test_importacao_isolada_por_empresa(moto, empresa_b, django_capture_on_commit_callbacks):
    _rodando(
        django_capture_on_commit_callbacks,
        _enviar,
        moto.api(),
        _xlsx(COLUNAS_VEICULO, _frota(1)),
    )
    connection.set_tenant(moto.empresa)
    assert Importacao.objects.count() == 1
    connection.set_tenant(empresa_b)
    assert Importacao.objects.count() == 0
    assert LinhaImportacao.objects.count() == 0
    connection.set_tenant(moto.empresa)


def test_importacao_lista_dos_ultimos_30_dias(moto, relogio, django_capture_on_commit_callbacks):
    api = moto.api()
    _rodando(django_capture_on_commit_callbacks, _enviar, api, _xlsx(COLUNAS_VEICULO, _frota(1)))
    assert len(api.get(URL).json()["resultados"]) == 1
    relogio.em("2026-11-07 10:00")
    assert not selectors.recentes().exists()


def test_cli_01_cpf_ja_cadastrado_sem_dizer_de_quem(moto, django_capture_on_commit_callbacks):
    from apps.clientes import services as clientes

    clientes.cadastrar(
        {
            "tipo": "pf",
            "nome": "Juliana Rocha",
            "documento": "05221833018",
            "celular": "85998123344",
        }
    )
    api = moto.api()
    imp = _rodando(
        django_capture_on_commit_callbacks,
        _enviar,
        api,
        _xlsx(COLUNAS_CLIENTE, [JULIANA]),
        "clientes-juliana.xlsx",
        "clientes",
    ).json()["id"]
    assert _linhas(api, imp)[0]["erros"] == [
        {"campo": "documento", "codigo": "CLI-01", "mensagem": "CPF já cadastrado"}
    ]
    # Nenhum evento guarda o nome do arquivo nem o conteúdo das linhas (ADR 0014).
    for evento in EventoAuditoria.objects.filter(objeto_tipo="importacao"):
        assert "clientes-juliana" not in str(evento.depois) + str(evento.antes)


def test_imp_02_arquivo_ilegivel_depois_de_aceito_falha_com_mensagem(
    moto, monkeypatch, django_capture_on_commit_callbacks
):
    api = moto.api()
    monkeypatch.setattr(arquivos.ProvedorFalso, "ler", lambda self, caminho: b"isto nao e zip")
    imp = _rodando(
        django_capture_on_commit_callbacks, _enviar, api, _xlsx(COLUNAS_VEICULO, _frota(1))
    ).json()["id"]
    corpo = api.get(f"{URL}{imp}/").json()
    assert corpo["status"] == "falhou"
    assert corpo["aviso"] == {
        "codigo": "validacao",
        "mensagem": "O arquivo não é uma planilha .xlsx.",
    }
